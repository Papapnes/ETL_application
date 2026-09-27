import io
import json
import zipfile

import pandas as pd
import pytest

from etl.demo import demo
from etl.engine import ETLError, bundle, default_recipe, extract, load_recipe, run


def test_demo_conservation_and_provenance():
    datasets, recipe = demo()
    result = run(datasets, recipe)
    assert result["report"]["input_rows"] == 12
    assert result["report"]["accepted_rows"] == 8
    assert result["report"]["rejected_rows"] == 4
    assert result["report"]["duplicate_rows"] == 1
    assert result["accepted"].iloc[0]["commande"] == "0001"
    assert result["accepted"].iloc[0]["montant"] == 1250.5
    assert result["rejected"]["_etl_record"].tolist() == [4, 6, 7, 9]
    assert datasets[0].frame.iloc[0]["client"] == " Atelier Nord "


@pytest.mark.parametrize("fmt", ["CSV", "Excel", "Parquet"])
def test_exports_roundtrip_and_quality_gate(fmt):
    datasets, recipe = demo()
    result = run(datasets, recipe)
    with pytest.raises(ETLError):
        bundle(result, fmt)
    with zipfile.ZipFile(io.BytesIO(bundle(result, fmt, True))) as z:
        report = json.loads(z.read("rapport.json"))
        assert report["partial_export_acknowledged"] is True
        assert report["accepted_rows"] == 8
        suffix = {"CSV": "csv", "Excel": "xlsx", "Parquet": "parquet"}[fmt]
        # Read exports directly: _etl_ provenance columns are deliberately reserved by import.
        content = io.BytesIO(z.read(f"donnees.{suffix}"))
        df = (
            pd.read_csv(content, sep=";", dtype={"commande": str})
            if fmt == "CSV"
            else pd.read_excel(content, dtype={"commande": str})
            if fmt == "Excel"
            else pd.read_parquet(content)
        )
        assert len(df) == 8
        assert df.iloc[0]["commande"] == "0001"
        assert "_etl_sha256" in df
        assert len(pd.read_csv(io.BytesIO(z.read("quarantaine.csv")), sep=";")) == 4


@pytest.mark.parametrize("content", [b"a,a\n1,2", b"a,b\n1,2,3", b"a,b\n1", b"_etl_source,b\n1,2"])
def test_bad_csv_is_not_silently_fixed(content):
    with pytest.raises(ETLError):
        extract("bad.csv", content, {"delimiter": ","})


def test_csv_quotes_newlines_encoding_and_leading_zeros():
    data = 'id;nom\n001;"Été\nParis"\n002;"A;B"'.encode("cp1252")
    df = extract("data.csv", data, {"encoding": "cp1252", "delimiter": ";"})[0].frame
    assert df.iloc[0].tolist() == ["001", "Été\nParis"]
    assert df.iloc[1]["nom"] == "A;B"


def test_nested_json_keeps_arrays_without_multiplying_rows():
    data = json.dumps({"data": {"items": [{"id": "001", "tags": [1, 2], "obj": {"a": 1}}]}}).encode()
    df = extract("data.json", data, {"json_path": "data.items"})[0].frame
    assert len(df) == 1
    assert json.loads(df.iloc[0]["tags"]) == [1, 2]


def test_jsonl():
    df = extract("data.jsonl", b'{"id":"001"}\n{"id":"002"}\n')[0].frame
    assert df["id"].tolist() == ["001", "002"]


@pytest.mark.parametrize("fmt", ["xlsx", "xls", "parquet"])
def test_binary_importers(fmt):
    buffer = io.BytesIO()
    data = pd.DataFrame({"id": ["001", "002"], "name": ["A", "B"]})
    if fmt == "xlsx":
        with pd.ExcelWriter(buffer) as writer:
            data.to_excel(writer, sheet_name="Clients", index=False)
            data.to_excel(writer, sheet_name="Autres", index=False)
    elif fmt == "xls":
        import xlwt

        book = xlwt.Workbook()
        sheet = book.add_sheet("Clients")
        for r, row in enumerate([list(data.columns)] + data.values.tolist()):
            for c, value in enumerate(row):
                sheet.write(r, c, value)
        book.save(buffer)
    else:
        data.to_parquet(buffer, index=False)
    tables = extract(f"data.{fmt}", buffer.getvalue())
    assert len(tables) == (2 if fmt == "xlsx" else 1)
    assert tables[0].frame["id"].tolist() == ["001", "002"]


def test_pdf_review_gate():
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    buffer = io.BytesIO()
    table = Table([["id", "name"], ["001", "Alpha"], ["002", "Beta"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, "black")]))
    SimpleDocTemplate(buffer).build([table])
    tables = extract("data.pdf", buffer.getvalue())
    assert len(tables) == 1
    recipe = default_recipe(list(tables[0].frame.columns))
    with pytest.raises(ETLError, match="PDF"):
        run(tables, recipe)
    assert len(run(tables, recipe, {tables[0].key})["accepted"]) == 2


def test_pdf_without_table_has_actionable_error():
    from reportlab.pdfgen.canvas import Canvas

    buffer = io.BytesIO()
    canvas = Canvas(buffer)
    canvas.drawString(50, 700, "No table")
    canvas.save()
    with pytest.raises(ETLError, match="Aucune table"):
        extract("text.pdf", buffer.getvalue())


def test_schema_drift_blocks_run_and_recipe_is_validated():
    sources = extract("data.csv", b"a,b\n1,2")
    recipe = default_recipe(["a", "c"])
    with pytest.raises(ETLError, match="Schéma"):
        run(sources, recipe)
    recipe = default_recipe(["a", "b"])
    recipe["rules"][1]["target"] = "a"
    with pytest.raises(ETLError, match="même nom"):
        load_recipe(json.dumps(recipe).encode())


def test_invalid_row_does_not_displace_valid_duplicate():
    sources = extract("data.csv", b"id,n\n001,bad\n001,10\n001,10")
    recipe = default_recipe(["id", "n"])
    recipe.update(deduplicate=True, duplicate_keys=["id"])
    recipe["rules"][1]["type"] = "Entier"
    result = run(sources, recipe)
    assert result["accepted"]["_etl_record"].tolist() == [2]
    assert result["report"]["duplicate_rows"] == 1


def test_spreadsheet_formula_protection():
    sources = extract("data.csv", b'name,n\n"=1+1",-2\n"  @SUM(1)",3')
    recipe = default_recipe(["name", "n"])
    recipe["rules"][1]["type"] = "Entier"
    result = run(sources, recipe)
    with zipfile.ZipFile(io.BytesIO(bundle(result, "Excel"))) as z:
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(z.read("donnees.xlsx"))).active
        assert sheet["A2"].data_type == "s"
        assert sheet["A2"].value == "'=1+1"
        assert sheet["B2"].value == -2


def test_batch_recipe_reuse_and_integer_precision():
    sources = extract("a.csv", b"id,n\n001,2\n002,9007199254740993\n003,2.5\n004,9007199254740991.1")
    recipe = default_recipe(["id", "n"])
    recipe["rules"][1]["type"] = "Entier"
    assert len(run(sources, recipe)["accepted"]) == 1
    more = extract("b.csv", b"id,n\n004,4")
    assert len(run(more, load_recipe(json.dumps(recipe).encode()))["accepted"]) == 1


def test_empty_rows_and_unknown_format():
    sources = extract("empty.csv", b"a,b\n", {"delimiter": ","})
    result = run(sources, default_recipe(["a", "b"]))
    assert result["report"]["input_rows"] == 0
    with pytest.raises(ETLError, match="Aucune ligne"):
        bundle(result, "CSV")
    with pytest.raises(ETLError, match="non pris"):
        extract("data.bin", b"anything")


def test_saved_recipe_on_next_batch_and_snapshot():
    from pathlib import Path

    root = Path(__file__).resolve().parents[1] / "examples"
    sources = extract("ventes_suivantes.csv", (root / "ventes_suivantes.csv").read_bytes())
    recipe = load_recipe((root / "recette_ventes.json").read_bytes())
    result = run(sources, recipe)
    assert len(result["accepted"]) == 3
    assert len(result["rejected"]) == 0
    recipe["name"] = "Modified later"
    assert result["recipe"]["name"] == "Préparation des ventes"
    assert bundle(result, "CSV")[:2] == b"PK"


def test_numeric_bounds_boolean_and_missing_markers():
    sources = extract("data.csv", b"n,active,name\n-1,oui,NULL\n10,non,Alpha\n20,perhaps,Beta")
    recipe = default_recipe(["n", "active", "name"])
    recipe["rules"][0].update(type="Entier", minimum="0", maximum="15")
    recipe["rules"][1]["type"] = "Booléen"
    recipe["rules"][2].update(required=True, fill="Unknown")
    recipe["null_markers"].append("NULL")
    result = run(sources, recipe)
    assert result["accepted"]["_etl_record"].tolist() == [2]
    assert not result["accepted"].iloc[0]["active"]
    assert "hors borne" in result["rejected"].iloc[0]["_etl_issues"]


def test_multi_source_counts_and_original_schema_order():
    sources = extract("one.csv", b"a,b\n001,x") + extract("two.csv", b"b,a\ny,002")
    result = run(sources, default_recipe(["a", "b"]))
    assert result["accepted"]["a"].tolist() == ["001", "002"]
    assert result["accepted"]["_etl_source"].tolist() == ["one.csv", "two.csv"]
